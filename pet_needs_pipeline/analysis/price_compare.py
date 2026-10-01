"""H14 가격 비교표 생성: 대안별 총 비용 (초기비용 + 월비용 × 기간).
가격은 2026-10 웹 검색 기준(출처·신뢰도는 '가격조사' 시트). 파란 글씨 셀은 가정값이라 팀이 바꾸면 총비용이 자동 재계산된다.
실행: python -m analysis.price_compare   → outputs/tables/price_compare.xlsx
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent.parent / "outputs/tables/price_compare.xlsx"
BLUE, HEAD = Font(color="0000FF"), PatternFill("solid", fgColor="DDE6F3")
WARN = PatternFill("solid", fgColor="FFF2CC")

wb = Workbook()
a = wb.active
a.title = "가정"
for r in [
    ("항목", "값", "설명"),
    ("비교 기간(개월)", 36, "총비용을 계산할 사용 기간 (로봇청소기 수명 3년 가정, 변경 가능)"),
    ("환율(원/USD)", 1400, "가정값. 퍼보 한국 정가를 못 찾아 미국 정가를 환산. 실제 환율로 바꿀 것"),
    ("펫시터 주당 방문 횟수", 3, "혼자 있는 평일 중 방문 횟수 가정"),
    ("펫시터 1회 요금(원)", 28000, "숨고 방문 펫시터 평균 28,000원 (최저 20,000 / 최고 40,000)"),
    ("우리 제품 하드웨어 추가금(원)", 0, "미정 — 간식 수납·배출 장치 원가(R4)가 나오면 입력. 0 = ThinQ 소프트웨어 업데이트만"),
    ("우리 제품 월 구독료(원)", 5000, "미정 — 설문(H12 Van Westendorp) 결과로 교체. 현재는 임시 가정"),
]:
    a.append(r)
for c in a["B"][1:]:
    c.font = BLUE
for c in (a["B6"], a["B7"], a["B3"]):
    c.fill = WARN

p = wb.create_sheet("가격조사")
p.append(["구분", "제품·항목", "가격(원)", "가격 기준", "출처", "신뢰도·비고"])
rows = [
    ("LG 로봇청소기", "코드제로 R5 오브제컬렉션 (LG전자 공식몰 판매가)", 570000, "정상가 90만원에서 36% 할인", "https://www.lge.co.kr/vacuum-cleaners/r585wka1", "검색 요약 인용. 공식 페이지는 접속 차단으로 직접 확인 못 함 → 팀이 최신가 재확인 필요"),
    ("LG 로봇청소기", "코드제로 R5 (타 판매처 회원할인 적용)", 460000, "판매가 69만원 − 회원할인 23만원", "https://plan.danawa.com/info/?nPlanSeq=7608", "2차 요약. 판매처 조건부 가격"),
    ("LG 로봇청소기", "코드제로 R5 정상가", 900000, "할인 전 판매가", "https://www.lge.co.kr/product/vacuum-cleaners/a720wa", "표기된 할인 전 판매가 (모델 일치 확인 필요)"),
    ("LG 로봇청소기", "코드제로 오브제컬렉션 R9 올인원타워 출하가", 1590000, "오브제 색상 159만 / 아이언그레이 149만", "https://www.lg.co.kr/media/release/25385", "보도자료 기반 출하가"),
    ("LG 로봇청소기", "코드제로 R9 (다나와 최저가, G마켓)", 1017270, "실구매가 사례 (판매가 1,374,000 / 최대혜택가 1,088,100)", "https://prod.danawa.com/info/?pcode=17978366", "시점·모델(RO965WB) 변동 큼. 재확인 필요"),
    ("경쟁 로봇청소기", "삼성 비스포크 AI 스팀 일반형 출고가(하한)", 1410000, "141만~159만원", "https://www.dailian.co.kr/news/view/1654746/", "경쟁사 참고용 (비교표에는 미포함)"),
    ("경쟁 로봇청소기", "에코백스 디봇 T30 프로 옴니", 599000, "판매 사례가", "https://www.smarttoday.co.kr/ko-kr/articles/74140", "프로모션 사례. 경쟁사 참고용"),
    ("일반 펫캠", "샤오미 MJSXJ05CM", 37500, "3.5만~4만원 중간값", "https://wikidocs.net/blog/@dognyang/24526/", "블로그 정보. 구독료 없음"),
    ("일반 펫캠", "TP-Link Tapo C210 (2K 팬틸트)", 45900, "판매가", "https://wikidocs.net/blog/@dognyang/24526/", "블로그 정보. 로컬 SD 저장, 구독 없이 사용"),
    ("일반 펫캠", "캠플러스 미지아 IPC-013", 65000, "6.5만원 전후", "https://wikidocs.net/blog/@dognyang/24526/", "블로그 정보"),
    ("간식 디스펜서 펫캠", "펫킷 얌쉐어 (1080P 카메라 AI 급식기)", 199000, "쿠팡 할인가", "https://www.coupang.com/vp/products/8999350932", "쿠팡 검색 요약. 간식 '투척' 방식인지는 미확인 (급식기)"),
    ("간식 디스펜서 펫캠", "IOT 스마트 자동급식기 (카메라 내장)", 96500, "쿠팡 할인가", "https://www.coupang.com/vp/products/8362062812", "쿠팡 검색 요약. 일반 급식기에 가까움"),
    ("간식 디스펜서 펫캠", "Furbo 360° 기기 (미국 정가 $199 × 환율)", "=199*가정!B3", "$199 MSRP (행사 시 $149~179)", "https://furbo.com/us/products/furbo-dog-cam-360", "한국 원화 정가는 못 찾음 → 환산값. 한국 판매가 재확인 필요"),
    ("간식 디스펜서 펫캠", "Furbo Nanny 구독 Basic (월, $6.99 × 환율)", "=6.99*가정!B3", "월 $6.99 (연 $83.88) / Premium $9.99", "https://thesmartsnout.com/2026/02/03/furbo-360-dog-camera-price-review-2026/", "구독 없이도 라이브뷰·수동 양방향·수동 간식 가능, 녹화·알림은 구독 필요"),
    ("펫시터", "방문 펫시터 1회 (숨고 평균)", 28000, "건당, 최저 20,000 / 최고 40,000", "https://soomgo.com/a.%EB%B6%80%EC%82%B0-%EC%82%AC%ED%95%98%EA%B5%AC/%ED%8E%AB-%EC%8B%9C%ED%84%B0-%EB%B0%A9%EB%AC%B8", "시간·지역별 상이"),
    ("펫시터", "와요 방문 펫시터 60분", 28500, "30분 19,500 / 120분 42,000", "https://apps.apple.com/us/app/%ED%8E%AB%EC%8B%9C%ED%84%B0-%EC%82%B0%EC%B1%85-%EB%B0%A9%EB%AC%B8%ED%83%81%EB%AC%98-%EC%BA%A3%EC%8B%9C%ED%84%B0-%EC%95%A0%EA%B2%AC%ED%98%B8%ED%85%94-%EB%B0%98%EB%A0%A4%EB%8F%99%EB%AC%BC-%ED%8E%AB%EB%B3%B4%ED%97%98/id1598224377?l=ko", "플랫폼 표기 시작가"),
]
for r in rows:
    p.append(list(r))
# 행 번호: 2 R5공식, 3 R5회원, 4 R5정상, 5 R9출하, 6 R9최저, 7 삼성, 8 에코백스, 9 샤오미, 10 Tapo, 11 미지아, 12 펫킷, 13 IOT, 14 Furbo, 15 Furbo구독, 16 숨고, 17 와요

t = wb.create_sheet("대안별 총비용")
t.append(["대안", "구성", "초기 비용(원)", "월 비용(원)", "총비용(기간, 원)", "월평균 환산(원)", "우리 제품 대비(원, 같은 R5/R9 기준끼리)", "비고"])
sc = [
    ("① 일반 로봇청소기 + 일반 펫캠", "LG R5 + Tapo C210", "=가격조사!C2+가격조사!C10", "=0", "구독 없음. 간식 투척·로봇 연동 없음"),
    ("② 로봇청소기 + 간식디스펜서 펫캠(국내)", "LG R5 + 펫킷 얌쉐어", "=가격조사!C2+가격조사!C12", "=0", "급식기형. 투척 여부 미확인"),
    ("③ 로봇청소기 + 퍼보(구독 포함)", "LG R5 + Furbo 360 + Nanny Basic", "=가격조사!C2+가격조사!C14", "=가격조사!C15", "한국 정가 미확인 → 환산값"),
    ("④ 프리미엄 로봇청소기 + 퍼보(구독 포함)", "LG R9(최저가) + Furbo 360 + Nanny Basic", "=가격조사!C6+가격조사!C14", "=가격조사!C15", "R9 가격은 변동 큼"),
    ("⑤ 펫시터 방문", "LG R5 + 방문 펫시터 (주 N회)", "=가격조사!C2", "=가정!B4*가정!B5*52/12", "펫시터 횟수·단가는 가정 시트에서 조정"),
    ("⑥ 우리 제품 (R5 기반)", "LG R5 + 펫케어 기능(하드웨어 추가금 + 월 구독)", "=가격조사!C2+가정!B6", "=가정!B7", "추가금·구독료는 미정(가정값)"),
    ("⑦ 우리 제품 (R9 기반)", "LG R9(최저가) + 펫케어 기능", "=가격조사!C6+가정!B6", "=가정!B7", "추가금·구독료는 미정(가정값)"),
]
for i, (n, comp, init, mon, note) in enumerate(sc, start=2):
    t.append([n, comp, init, mon, f"=C{i}+D{i}*가정!$B$2", f"=E{i}/가정!$B$2",
              {2: "=E2-E$7", 3: "=E3-E$7", 4: "=E4-E$7", 5: "=E5-E$8", 6: "=E6-E$7", 7: "", 8: ""}[i], note])
t.append([])
t.append(["주의", "총비용 = 초기 비용 + 월 비용 × 비교 기간. 대안①~④는 '우리 제품'과 같은 R5/R9 기준끼리 비교(④·⑦은 R9 기준, ①~③·⑤·⑥은 R5 기준)."])
t.append(["", "가격은 검색 요약 기반이고 공식몰·다나와는 접속 차단으로 직접 확인하지 못했다. 발표 전에 최신가·구독 조건을 재확인할 것."])
t.append(["", "비용만 비교한 표이다. '더 높은 가치'(H14)는 비용 차이만으로 판단할 수 없고, 설문 선호(우리 제품 선호 > 대안 선호)와 함께 제시해야 한다."])

for ws in (a, p, t):
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), HEAD
    ws.freeze_panes = "A2"
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 28 if col != 5 else 40
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
for ws, cols in ((p, "C"), (t, "CDEFG")):
    for col in cols:
        for c in ws[col][1:]:
            c.number_format = "#,##0"
for c in a["B"][3:]:
    c.number_format = "#,##0"
a.column_dimensions["C"].width = 70
wb.save(OUT)
print("saved", OUT)
